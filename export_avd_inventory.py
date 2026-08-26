#!/usr/bin/env python3
"""Export Azure Virtual Desktop inventory for one subscription to a multi-sheet Excel workbook.

Sheets:
  1. HostPools_SessionHosts  – host pool details + each associated session host
  2. AppGroups_Access        – host pool, type, application groups, assignments, apps
  3. HostPools               – one row per host pool
  4. SessionHosts            – one row per session host
  5. ApplicationGroups       – one row per application group
  6. Assignments             – who is assigned to each application group
  7. Applications            – published apps/desktops per application group
  8. Workspaces              – workspaces and linked application groups
  9. Summary                 – counts

Authentication uses Azure Identity (az login, environment, managed identity, or device code).

Required Azure permissions:
  - Desktop Virtualization Reader (or higher) on the subscription
  - Microsoft.Authorization/roleAssignments/read
Optional Graph permission (Directory.Read.All) to resolve user/group names.
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
from collections import defaultdict
from datetime import datetime, timezone
from typing import Any, Iterable

try:
    import requests
    from azure.core.exceptions import HttpResponseError
    from azure.identity import (
        AzureCliCredential,
        DefaultAzureCredential,
        DeviceCodeCredential,
        InteractiveBrowserCredential,
    )
    from azure.mgmt.authorization import AuthorizationManagementClient
    from azure.mgmt.desktopvirtualization import DesktopVirtualizationMgmtClient
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet
except ModuleNotFoundError as exc:
    missing = getattr(exc, "name", None) or str(exc)
    print(
        f"Missing Python package: {missing}\n"
        "Install dependencies into THIS same Python interpreter:\n"
        f"  {sys.executable} -m pip install -r requirements.txt\n"
        "If you use a venv, activate it first, then run the script with that venv's python.",
        file=sys.stderr,
    )
    raise SystemExit(1) from exc

try:
    from azure.mgmt.compute import ComputeManagementClient
except ModuleNotFoundError:
    ComputeManagementClient = None  # type: ignore[misc, assignment]

LOGGER = logging.getLogger("avd_inventory")
SCRIPT_VERSION = "2026.08.26-1"

HOST_POOL_ID_RE = re.compile(
    r"/subscriptions/[^/]+/resourcegroups/([^/]+)/providers/"
    r"microsoft\.desktopvirtualization/hostpools/([^/]+)",
    re.IGNORECASE,
)
RESOURCE_GROUP_RE = re.compile(r"/resourcegroups/([^/]+)/", re.IGNORECASE)
COMPUTE_VM_RE = re.compile(
    r"/resourcegroups/([^/]+)/providers/microsoft\.compute/virtualmachines/([^/]+)$",
    re.IGNORECASE,
)
COMPUTE_VMSS_RE = re.compile(
    r"/resourcegroups/([^/]+)/providers/microsoft\.compute/"
    r"virtualmachinescalesets/([^/]+)/virtualmachines/([^/]+)$",
    re.IGNORECASE,
)

GRAPH_BASE = "https://graph.microsoft.com/v1.0"


def enum_value(value: Any) -> str:
    if value is None:
        return ""
    return str(getattr(value, "value", value))


def model_attr(obj: Any, *names: str, default: Any = "") -> Any:
    """Read the first present SDK attribute. Models differ slightly by API version."""
    for name in names:
        if hasattr(obj, name):
            value = getattr(obj, name)
            if value is not None:
                return value
    return default


def iso_dt(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def parse_resource_group(resource_id: str | None) -> str:
    if not resource_id:
        return ""
    match = RESOURCE_GROUP_RE.search(resource_id)
    return match.group(1) if match else ""


def parse_host_pool_from_arm(host_pool_arm_path: str | None) -> tuple[str, str]:
    if not host_pool_arm_path:
        return "", ""
    match = HOST_POOL_ID_RE.search(host_pool_arm_path)
    if not match:
        return "", host_pool_arm_path.rstrip("/").split("/")[-1]
    return match.group(1), match.group(2)


def join_unique(values: Iterable[Any]) -> str:
    seen: list[str] = []
    for value in values:
        text = str(value).strip() if value is not None else ""
        if text and text not in seen:
            seen.append(text)
    return "; ".join(seen)


def short_session_host_name(name: str | None) -> str:
    """Azure returns session hosts as hostPoolName/sessionHostName; keep only the host."""
    if not name:
        return ""
    return str(name).replace("\\", "/").rstrip("/").split("/")[-1]


def get_credential(auth_mode: str, tenant_id: str | None):
    kwargs: dict[str, Any] = {}
    if tenant_id:
        kwargs["tenant_id"] = tenant_id

    if auth_mode == "device":
        LOGGER.info("Using device-code authentication. Complete the prompt in a browser.")
        return DeviceCodeCredential(**kwargs)
    if auth_mode == "interactive":
        LOGGER.info("Using interactive browser authentication.")
        return InteractiveBrowserCredential(**kwargs)
    if auth_mode == "cli":
        LOGGER.info("Using Azure CLI credential (run 'az login' first).")
        return AzureCliCredential(**kwargs)

    LOGGER.info("Using DefaultAzureCredential (CLI, environment, managed identity, browser).")
    return DefaultAzureCredential(exclude_interactive_browser_credential=False, **kwargs)


class GraphDirectory:
    """Resolve Azure AD object IDs to display names via Microsoft Graph."""

    def __init__(self, credential):
        self._credential = credential
        self._cache: dict[str, dict[str, str]] = {}
        self._enabled = True

    def _token(self) -> str | None:
        try:
            return self._credential.get_token("https://graph.microsoft.com/.default").token
        except Exception as exc:  # noqa: BLE001
            LOGGER.warning("Microsoft Graph token unavailable (%s). Principal names will be blank.", exc)
            self._enabled = False
            return None

    def resolve_many(self, object_ids: Iterable[str]) -> None:
        pending = [oid for oid in dict.fromkeys(object_ids) if oid and oid not in self._cache]
        if not pending or not self._enabled:
            return

        token = self._token()
        if not token:
            return

        headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
        for offset in range(0, len(pending), 1000):
            chunk = pending[offset : offset + 1000]
            try:
                response = requests.post(
                    f"{GRAPH_BASE}/directoryObjects/getByIds",
                    headers=headers,
                    json={"ids": chunk, "types": []},
                    timeout=60,
                )
                if response.status_code >= 400:
                    LOGGER.warning(
                        "Graph getByIds failed (%s): %s",
                        response.status_code,
                        response.text[:300],
                    )
                    self._enabled = False
                    return
                for item in response.json().get("value", []):
                    self._cache[item.get("id", "")] = self._principal_from_graph(item)
            except requests.RequestException as exc:
                LOGGER.warning("Graph lookup failed: %s", exc)
                self._enabled = False
                return

            for object_id in chunk:
                self._cache.setdefault(
                    object_id,
                    {
                        "principal_id": object_id,
                        "principal_type": "Unknown",
                        "display_name": "",
                        "user_principal_name": "",
                        "mail": "",
                    },
                )

    def get(self, object_id: str) -> dict[str, str]:
        if not object_id:
            return {
                "principal_id": "",
                "principal_type": "",
                "display_name": "",
                "user_principal_name": "",
                "mail": "",
            }
        self.resolve_many([object_id])
        return self._cache.get(
            object_id,
            {
                "principal_id": object_id,
                "principal_type": "Unknown",
                "display_name": "",
                "user_principal_name": "",
                "mail": "",
            },
        )

    @staticmethod
    def _principal_from_graph(item: dict[str, Any]) -> dict[str, str]:
        odata_type = (item.get("@odata.type") or "").lower()
        if "user" in odata_type:
            principal_type = "User"
        elif "group" in odata_type:
            principal_type = "Group"
        elif "serviceprincipal" in odata_type:
            principal_type = "ServicePrincipal"
        else:
            principal_type = odata_type.replace("#microsoft.graph.", "") or "DirectoryObject"

        return {
            "principal_id": item.get("id") or "",
            "principal_type": principal_type,
            "display_name": item.get("displayName") or item.get("appDisplayName") or "",
            "user_principal_name": item.get("userPrincipalName") or item.get("appId") or "",
            "mail": item.get("mail") or "",
        }


class AvdInventoryCollector:
    def __init__(self, credential, subscription_id: str, include_inherited_assignments: bool):
        self.subscription_id = subscription_id
        self.include_inherited_assignments = include_inherited_assignments
        self.avd = DesktopVirtualizationMgmtClient(credential, subscription_id)
        self.auth = AuthorizationManagementClient(credential, subscription_id)
        self.graph = GraphDirectory(credential)
        self._role_name_cache: dict[str, str] = {}
        self._os_cache: dict[str, dict[str, str]] = {}
        self.compute = None
        if ComputeManagementClient is not None:
            self.compute = ComputeManagementClient(credential, subscription_id)

    def collect(self) -> dict[str, list[dict[str, Any]]]:
        LOGGER.info("Listing host pools in subscription %s", self.subscription_id)
        host_pools = list(self.avd.host_pools.list())
        LOGGER.info("Found %s host pool(s)", len(host_pools))

        LOGGER.info("Listing application groups")
        app_groups = list(self.avd.application_groups.list_by_subscription())
        LOGGER.info("Found %s application group(s)", len(app_groups))

        LOGGER.info("Listing workspaces")
        try:
            workspaces = list(self.avd.workspaces.list_by_subscription())
        except HttpResponseError as exc:
            LOGGER.warning("Unable to list workspaces: %s", exc)
            workspaces = []
        LOGGER.info("Found %s workspace(s)", len(workspaces))

        host_pool_rows: list[dict[str, Any]] = []
        session_host_rows: list[dict[str, Any]] = []
        combined_host_rows: list[dict[str, Any]] = []

        app_groups_by_pool: dict[str, list[Any]] = defaultdict(list)
        for app_group in app_groups:
            pool_id = (app_group.host_pool_arm_path or "").rstrip("/").lower()
            app_groups_by_pool[pool_id].append(app_group)

        workspace_rows = self._workspace_rows(workspaces)
        workspace_by_app_group: dict[str, list[str]] = defaultdict(list)
        for workspace in workspaces:
            for reference in workspace.application_group_references or []:
                workspace_by_app_group[reference.rstrip("/").lower()].append(workspace.name or "")

        for host_pool in host_pools:
            pool_row = self._host_pool_row(host_pool)
            rg = parse_resource_group(host_pool.id)
            LOGGER.info("Listing session hosts for %s", host_pool.name)
            try:
                session_hosts = list(self.avd.session_hosts.list(rg, host_pool.name))
            except HttpResponseError as exc:
                LOGGER.warning("Unable to list session hosts for %s: %s", host_pool.name, exc)
                session_hosts = []

            pool_row["SessionHostCount"] = len(session_hosts)
            pool_row["ApplicationGroupCount"] = len(app_groups_by_pool.get((host_pool.id or "").rstrip("/").lower(), []))
            host_pool_rows.append(pool_row)

            if not session_hosts:
                combined_host_rows.append(
                    {
                        **pool_row,
                        "SessionHostName": "",
                        "SessionHostStatus": "",
                        "SessionHostAllowNewSession": "",
                        "Sessions": "",
                        "AssignedUser": "",
                        "LastHeartBeat": "",
                        "AgentVersion": "",
                        "OSName": "",
                        "OSVersion": "",
                        "OSType": "",
                        "SxSStackVersion": "",
                        "SessionHostResourceId": "",
                        "UpdateState": "",
                        "UpdateErrorMessage": "",
                    }
                )
                continue

            for session_host in session_hosts:
                host_row = self._session_host_row(host_pool, session_host)
                session_host_rows.append(host_row)
                combined_host_rows.append({**pool_row, **host_row})

        app_group_rows: list[dict[str, Any]] = []
        assignment_rows: list[dict[str, Any]] = []
        application_rows: list[dict[str, Any]] = []
        access_rows: list[dict[str, Any]] = []

        principal_ids: list[str] = []
        pending_assignments: list[tuple[Any, Any, list[Any]]] = []

        pool_by_id = {(hp.id or "").rstrip("/").lower(): hp for hp in host_pools}

        for app_group in app_groups:
            host_pool = pool_by_id.get((app_group.host_pool_arm_path or "").rstrip("/").lower())
            assignments = self._list_assignments(app_group.id)
            pending_assignments.append((app_group, host_pool, assignments))
            principal_ids.extend(a.principal_id for a in assignments if a.principal_id)

        self.graph.resolve_many(principal_ids)

        for app_group, host_pool, assignments in pending_assignments:
            rg = parse_resource_group(app_group.id)
            applications = self._list_applications(rg, app_group)
            assignment_summaries = [self._assignment_summary(a) for a in assignments]
            assigned_to = join_unique(
                s["AssignedTo"] for s in assignment_summaries if s["AssignedTo"]
            ) or join_unique(s["PrincipalId"] for s in assignment_summaries)

            app_group_row = self._app_group_row(
                host_pool,
                app_group,
                applications,
                assignment_summaries,
                workspace_by_app_group.get((app_group.id or "").rstrip("/").lower(), []),
            )
            app_group_rows.append(app_group_row)

            for summary in assignment_summaries:
                assignment_rows.append(
                    {
                        "HostPoolName": app_group_row["HostPoolName"],
                        "HostPoolType": app_group_row["HostPoolType"],
                        "ApplicationGroupName": app_group.name or "",
                        "ApplicationGroupType": enum_value(app_group.application_group_type),
                        "ApplicationGroupResourceGroup": rg,
                        **summary,
                    }
                )

            if not applications:
                application_rows.append(
                    {
                        "HostPoolName": app_group_row["HostPoolName"],
                        "HostPoolType": app_group_row["HostPoolType"],
                        "ApplicationGroupName": app_group.name or "",
                        "ApplicationGroupType": enum_value(app_group.application_group_type),
                        "ApplicationName": "",
                        "ApplicationFriendlyName": "",
                        "ApplicationType": "",
                        "FilePath": "",
                        "CommandLineSetting": "",
                        "CommandLineArguments": "",
                        "ShowInPortal": "",
                        "IconPath": "",
                        "AssignedTo": assigned_to,
                    }
                )
            else:
                for application in applications:
                    app_row = self._application_row(app_group_row, app_group, application, assigned_to)
                    application_rows.append(app_row)

            if not applications and not assignment_summaries:
                access_rows.append(
                    {
                        **{k: app_group_row[k] for k in (
                            "HostPoolName",
                            "HostPoolType",
                            "HostPoolResourceGroup",
                            "ApplicationGroupName",
                            "ApplicationGroupType",
                            "ApplicationGroupFriendlyName",
                            "Workspaces",
                        )},
                        "ApplicationName": "",
                        "ApplicationFriendlyName": "",
                        "ApplicationFilePath": "",
                        "PrincipalType": "",
                        "AssignedTo": "",
                        "UserPrincipalName": "",
                        "Mail": "",
                        "Role": "",
                        "AssignmentScope": "",
                        "AssignmentLevel": "",
                    }
                )
            elif not assignment_summaries:
                for application in applications:
                    access_rows.append(
                        {
                            **{k: app_group_row[k] for k in (
                                "HostPoolName",
                                "HostPoolType",
                                "HostPoolResourceGroup",
                                "ApplicationGroupName",
                                "ApplicationGroupType",
                                "ApplicationGroupFriendlyName",
                                "Workspaces",
                            )},
                            "ApplicationName": application.get("ApplicationName", ""),
                            "ApplicationFriendlyName": application.get("ApplicationFriendlyName", ""),
                            "ApplicationFilePath": application.get("FilePath", ""),
                            "PrincipalType": "",
                            "AssignedTo": "",
                            "UserPrincipalName": "",
                            "Mail": "",
                            "Role": "",
                            "AssignmentScope": "",
                            "AssignmentLevel": "",
                        }
                    )
            elif not applications:
                for summary in assignment_summaries:
                    access_rows.append(
                        {
                            **{k: app_group_row[k] for k in (
                                "HostPoolName",
                                "HostPoolType",
                                "HostPoolResourceGroup",
                                "ApplicationGroupName",
                                "ApplicationGroupType",
                                "ApplicationGroupFriendlyName",
                                "Workspaces",
                            )},
                            "ApplicationName": "",
                            "ApplicationFriendlyName": "",
                            "ApplicationFilePath": "",
                            "PrincipalType": summary["PrincipalType"],
                            "AssignedTo": summary["AssignedTo"],
                            "UserPrincipalName": summary["UserPrincipalName"],
                            "Mail": summary["Mail"],
                            "Role": summary["Role"],
                            "AssignmentScope": summary["AssignmentScope"],
                            "AssignmentLevel": summary["AssignmentLevel"],
                        }
                    )
            else:
                for application in applications:
                    for summary in assignment_summaries:
                        access_rows.append(
                            {
                                **{k: app_group_row[k] for k in (
                                    "HostPoolName",
                                    "HostPoolType",
                                    "HostPoolResourceGroup",
                                    "ApplicationGroupName",
                                    "ApplicationGroupType",
                                    "ApplicationGroupFriendlyName",
                                    "Workspaces",
                                )},
                                "ApplicationName": application.get("ApplicationName", ""),
                                "ApplicationFriendlyName": application.get("ApplicationFriendlyName", ""),
                                "ApplicationFilePath": application.get("FilePath", ""),
                                "PrincipalType": summary["PrincipalType"],
                                "AssignedTo": summary["AssignedTo"],
                                "UserPrincipalName": summary["UserPrincipalName"],
                                "Mail": summary["Mail"],
                                "Role": summary["Role"],
                                "AssignmentScope": summary["AssignmentScope"],
                                "AssignmentLevel": summary["AssignmentLevel"],
                            }
                        )

        summary_rows = [
            {"Metric": "SubscriptionId", "Value": self.subscription_id},
            {"Metric": "HostPools", "Value": len(host_pool_rows)},
            {"Metric": "SessionHosts", "Value": len(session_host_rows)},
            {"Metric": "ApplicationGroups", "Value": len(app_group_rows)},
            {"Metric": "Assignments", "Value": len(assignment_rows)},
            {"Metric": "Applications", "Value": sum(1 for r in application_rows if r.get("ApplicationName"))},
            {"Metric": "Workspaces", "Value": len(workspace_rows)},
        ]

        return {
            "HostPools_SessionHosts": combined_host_rows,
            "AppGroups_Access": access_rows,
            "HostPools": host_pool_rows,
            "SessionHosts": session_host_rows,
            "ApplicationGroups": app_group_rows,
            "Assignments": assignment_rows,
            "Applications": application_rows,
            "Workspaces": workspace_rows,
            "Summary": summary_rows,
        }

    def _host_pool_row(self, host_pool) -> dict[str, Any]:
        agent_update = host_pool.agent_update
        return {
            "HostPoolName": host_pool.name or "",
            "HostPoolType": enum_value(host_pool.host_pool_type),
            "HostPoolFriendlyName": host_pool.friendly_name or "",
            "HostPoolDescription": host_pool.description or "",
            "HostPoolResourceGroup": parse_resource_group(host_pool.id),
            "Location": host_pool.location or "",
            "LoadBalancerType": enum_value(host_pool.load_balancer_type),
            "MaxSessionLimit": host_pool.max_session_limit if host_pool.max_session_limit is not None else "",
            "PreferredAppGroupType": enum_value(host_pool.preferred_app_group_type),
            "PersonalDesktopAssignmentType": enum_value(host_pool.personal_desktop_assignment_type),
            "StartVMOnConnect": host_pool.start_vm_on_connect,
            "ValidationEnvironment": host_pool.validation_environment,
            "CustomRdpProperty": host_pool.custom_rdp_property or "",
            "VMTemplate": host_pool.vm_template or "",
            "Ring": host_pool.ring if host_pool.ring is not None else "",
            "HostPoolKind": enum_value(model_attr(host_pool, "kind", "host_pool_kind", default="")),
            "PublicNetworkAccess": enum_value(getattr(host_pool, "public_network_access", None)),
            "AgentUpdateType": enum_value(getattr(agent_update, "type", None)) if agent_update else "",
            "AgentUseSessionHostLocalTime": getattr(agent_update, "use_session_host_local_time", "") if agent_update else "",
            "Tags": join_unique(f"{k}={v}" for k, v in (host_pool.tags or {}).items()),
            "HostPoolId": host_pool.id or "",
        }

    def _session_host_row(self, host_pool, session_host) -> dict[str, Any]:
        assigned_user = session_host.assigned_user or ""
        if assigned_user:
            principal = self.graph.get(assigned_user) if re.fullmatch(r"[0-9a-fA-F-]{36}", assigned_user) else None
            if principal and principal.get("display_name"):
                assigned_user = f"{principal['display_name']} ({principal.get('user_principal_name') or assigned_user})"

        os_info = self._session_host_os(session_host)
        return {
            "HostPoolName": host_pool.name or "",
            "HostPoolType": enum_value(host_pool.host_pool_type),
            "SessionHostName": short_session_host_name(session_host.name),
            "OSName": os_info["OSName"],
            "OSVersion": os_info["OSVersion"],
            "OSType": os_info["OSType"],
            "SessionHostStatus": enum_value(session_host.status),
            "SessionHostAllowNewSession": session_host.allow_new_session,
            "Sessions": session_host.sessions if session_host.sessions is not None else "",
            "AssignedUser": assigned_user,
            "LastHeartBeat": iso_dt(session_host.last_heart_beat),
            "AgentVersion": session_host.agent_version or "",
            "SxSStackVersion": session_host.sx_s_stack_version or "",
            "SessionHostResourceId": session_host.resource_id or "",
            "UpdateState": enum_value(session_host.update_state),
            "UpdateErrorMessage": session_host.update_error_message or "",
        }

    def _session_host_os(self, session_host) -> dict[str, str]:
        avd_version = model_attr(session_host, "os_version")
        vm_info = self._vm_os_info(model_attr(session_host, "resource_id"))
        return {
            "OSName": vm_info.get("os_name") or "",
            "OSVersion": vm_info.get("os_version") or avd_version or "",
            "OSType": vm_info.get("os_type") or "",
        }

    def _vm_os_info(self, resource_id: str) -> dict[str, str]:
        empty = {"os_name": "", "os_version": "", "os_type": ""}
        if not resource_id or self.compute is None:
            return empty
        cache_key = resource_id.rstrip("/").lower()
        if cache_key in self._os_cache:
            return self._os_cache[cache_key]

        info = dict(empty)
        try:
            vmss_match = COMPUTE_VMSS_RE.search(resource_id)
            vm_match = COMPUTE_VM_RE.search(resource_id)
            resource = None
            if vmss_match:
                resource = self.compute.virtual_machine_scale_set_vms.get(
                    vmss_match.group(1),
                    vmss_match.group(2),
                    vmss_match.group(3),
                    expand="instanceView",
                )
            elif vm_match:
                resource = self.compute.virtual_machines.get(
                    vm_match.group(1),
                    vm_match.group(2),
                    expand="instanceView",
                )
            if resource is not None:
                instance_view = getattr(resource, "instance_view", None)
                info["os_name"] = model_attr(instance_view, "os_name") if instance_view else ""
                info["os_version"] = model_attr(instance_view, "os_version") if instance_view else ""
                storage = getattr(resource, "storage_profile", None)
                os_disk = getattr(storage, "os_disk", None) if storage else None
                info["os_type"] = enum_value(model_attr(os_disk, "os_type", default=None)) if os_disk else ""
                image = getattr(storage, "image_reference", None) if storage else None
                if not info["os_name"] and image:
                    offer = model_attr(image, "offer")
                    sku = model_attr(image, "sku")
                    info["os_name"] = " ".join(part for part in (offer, sku) if part)
                if not info["os_version"] and image:
                    info["os_version"] = model_attr(image, "exact_version") or model_attr(image, "version")
        except HttpResponseError as exc:
            LOGGER.debug("Unable to read VM OS for %s: %s", resource_id, exc)

        self._os_cache[cache_key] = info
        return info

    def _app_group_row(
        self,
        host_pool,
        app_group,
        applications: list[dict[str, Any]],
        assignment_summaries: list[dict[str, Any]],
        workspaces: list[str],
    ) -> dict[str, Any]:
        pool_name = host_pool.name if host_pool else parse_host_pool_from_arm(app_group.host_pool_arm_path)[1]
        pool_type = enum_value(host_pool.host_pool_type) if host_pool else ""
        pool_rg = parse_resource_group(host_pool.id) if host_pool else parse_host_pool_from_arm(app_group.host_pool_arm_path)[0]
        assigned_to = join_unique(s["AssignedTo"] or s["PrincipalId"] for s in assignment_summaries)
        app_names = join_unique(a.get("ApplicationName") for a in applications if a.get("ApplicationName"))
        return {
            "HostPoolName": pool_name,
            "HostPoolType": pool_type,
            "HostPoolResourceGroup": pool_rg,
            "ApplicationGroupName": app_group.name or "",
            "ApplicationGroupType": enum_value(app_group.application_group_type),
            "ApplicationGroupFriendlyName": app_group.friendly_name or "",
            "ApplicationGroupDescription": app_group.description or "",
            "ApplicationGroupResourceGroup": parse_resource_group(app_group.id),
            "Location": app_group.location or "",
            "Workspaces": join_unique(workspaces),
            "ApplicationCount": len([a for a in applications if a.get("ApplicationName")]),
            "Applications": app_names,
            "AssignmentCount": len(assignment_summaries),
            "AssignedTo": assigned_to,
            "ApplicationGroupId": app_group.id or "",
            "HostPoolArmPath": app_group.host_pool_arm_path or "",
        }

    def _application_row(self, app_group_row, app_group, application: dict[str, Any], assigned_to: str) -> dict[str, Any]:
        return {
            "HostPoolName": app_group_row["HostPoolName"],
            "HostPoolType": app_group_row["HostPoolType"],
            "ApplicationGroupName": app_group.name or "",
            "ApplicationGroupType": enum_value(app_group.application_group_type),
            "ApplicationName": application.get("ApplicationName", ""),
            "ApplicationFriendlyName": application.get("ApplicationFriendlyName", ""),
            "ApplicationType": application.get("ApplicationType", ""),
            "FilePath": application.get("FilePath", ""),
            "CommandLineSetting": application.get("CommandLineSetting", ""),
            "CommandLineArguments": application.get("CommandLineArguments", ""),
            "ShowInPortal": application.get("ShowInPortal", ""),
            "IconPath": application.get("IconPath", ""),
            "AssignedTo": assigned_to,
        }

    def _list_applications(self, resource_group: str, app_group) -> list[dict[str, Any]]:
        apps: list[dict[str, Any]] = []
        group_type = enum_value(app_group.application_group_type).lower()
        try:
            if group_type == "desktop":
                desktops = list(self.avd.desktops.list(resource_group, app_group.name))
                for desktop in desktops:
                    apps.append(
                        {
                            "ApplicationName": desktop.name or "SessionDesktop",
                            "ApplicationFriendlyName": desktop.friendly_name or "SessionDesktop",
                            "ApplicationType": "Desktop",
                            "FilePath": "",
                            "CommandLineSetting": "",
                            "CommandLineArguments": "",
                            "ShowInPortal": model_attr(desktop, "show_in_portal", "show_in_feed"),
                            "IconPath": model_attr(desktop, "icon_path"),
                        }
                    )
            published = list(self.avd.applications.list(resource_group, app_group.name))
            for application in published:
                apps.append(
                    {
                        "ApplicationName": application.name or "",
                        "ApplicationFriendlyName": application.friendly_name or application.name or "",
                        "ApplicationType": enum_value(model_attr(application, "application_type", default=None))
                        or "RemoteApp",
                        "FilePath": model_attr(application, "file_path"),
                        "CommandLineSetting": enum_value(model_attr(application, "command_line_setting", default=None)),
                        "CommandLineArguments": model_attr(application, "command_line_arguments"),
                        "ShowInPortal": model_attr(application, "show_in_portal"),
                        "IconPath": model_attr(application, "icon_path"),
                    }
                )
        except (HttpResponseError, AttributeError) as exc:
            LOGGER.warning("Unable to list applications for %s: %s", app_group.name, exc)
        return apps

    def _list_assignments(self, scope: str) -> list[Any]:
        if not scope:
            return []
        try:
            assignments = list(self.auth.role_assignments.list_for_scope(scope))
        except HttpResponseError as exc:
            LOGGER.warning("Unable to list role assignments for %s: %s", scope, exc)
            return []

        if self.include_inherited_assignments:
            return assignments

        normalized = scope.rstrip("/").lower()
        return [a for a in assignments if (a.scope or "").rstrip("/").lower() == normalized]

    def _role_name(self, role_definition_id: str | None) -> str:
        if not role_definition_id:
            return ""
        if role_definition_id in self._role_name_cache:
            return self._role_name_cache[role_definition_id]
        name = role_definition_id.rstrip("/").split("/")[-1]
        try:
            definition = self.auth.role_definitions.get_by_id(role_definition_id)
            name = definition.role_name or name
        except HttpResponseError:
            pass
        self._role_name_cache[role_definition_id] = name
        return name

    def _assignment_summary(self, assignment) -> dict[str, Any]:
        principal = self.graph.get(assignment.principal_id or "")
        assigned_to = principal.get("display_name") or assignment.principal_id or ""
        upn = principal.get("user_principal_name") or ""
        if upn and assigned_to and upn != assigned_to:
            assigned_display = f"{assigned_to} ({upn})"
        else:
            assigned_display = assigned_to or upn
        scope = assignment.scope or ""
        return {
            "PrincipalId": assignment.principal_id or "",
            "PrincipalType": principal.get("principal_type") or enum_value(assignment.principal_type),
            "AssignedTo": assigned_display,
            "UserPrincipalName": upn,
            "Mail": principal.get("mail") or "",
            "Role": self._role_name(assignment.role_definition_id),
            "AssignmentScope": scope,
            "AssignmentLevel": self._assignment_level(scope),
            "RoleAssignmentId": assignment.id or "",
        }

    @staticmethod
    def _assignment_level(scope: str) -> str:
        lower = (scope or "").lower()
        if "/applicationgroups/" in lower:
            return "ApplicationGroup"
        if "/resourcegroups/" in lower:
            return "ResourceGroup"
        if "/subscriptions/" in lower:
            return "Subscription"
        return "Other"

    def _workspace_rows(self, workspaces) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for workspace in workspaces:
            references = workspace.application_group_references or []
            if not references:
                rows.append(
                    {
                        "WorkspaceName": workspace.name or "",
                        "FriendlyName": workspace.friendly_name or "",
                        "Description": workspace.description or "",
                        "ResourceGroup": parse_resource_group(workspace.id),
                        "Location": workspace.location or "",
                        "ApplicationGroupName": "",
                        "ApplicationGroupId": "",
                        "WorkspaceId": workspace.id or "",
                    }
                )
                continue
            for reference in references:
                rows.append(
                    {
                        "WorkspaceName": workspace.name or "",
                        "FriendlyName": workspace.friendly_name or "",
                        "Description": workspace.description or "",
                        "ResourceGroup": parse_resource_group(workspace.id),
                        "Location": workspace.location or "",
                        "ApplicationGroupName": reference.rstrip("/").split("/")[-1],
                        "ApplicationGroupId": reference,
                        "WorkspaceId": workspace.id or "",
                    }
                )
        return rows


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL = PatternFill("solid", fgColor="D6EAF8")
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
WRAP = Alignment(vertical="top", wrap_text=True)


def write_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    sheet: Worksheet = workbook.create_sheet(title)
    sheet.append(columns)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        sheet.append([_excel_value(row.get(column, "")) for column in columns])

    for row_idx, excel_row in enumerate(sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=len(columns)), start=2):
        for cell in excel_row:
            cell.alignment = WRAP
            cell.border = THIN
        if row_idx % 2 == 0:
            for cell in excel_row:
                cell.fill = ALT_FILL

    sheet.freeze_panes = "A2"
    last_row = max(sheet.max_row, 1)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"

    for index, column in enumerate(columns, start=1):
        max_len = len(column)
        for cell in sheet.iter_cols(min_col=index, max_col=index, min_row=2, max_row=min(sheet.max_row, 200), values_only=True):
            for value in cell:
                if value is None:
                    continue
                max_len = max(max_len, min(len(str(value)), 60))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(12, max_len + 2), 48)

    if sheet.max_row == 1:
        sheet.append([""] * len(columns))


def _excel_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    return value


SHEET_COLUMNS = {
    "HostPools_SessionHosts": [
        "HostPoolName",
        "HostPoolType",
        "HostPoolFriendlyName",
        "HostPoolResourceGroup",
        "Location",
        "LoadBalancerType",
        "MaxSessionLimit",
        "PreferredAppGroupType",
        "PersonalDesktopAssignmentType",
        "StartVMOnConnect",
        "ValidationEnvironment",
        "SessionHostCount",
        "ApplicationGroupCount",
        "SessionHostName",
        "OSName",
        "OSVersion",
        "OSType",
        "SessionHostStatus",
        "SessionHostAllowNewSession",
        "Sessions",
        "AssignedUser",
        "LastHeartBeat",
        "AgentVersion",
        "SxSStackVersion",
        "SessionHostResourceId",
        "UpdateState",
        "UpdateErrorMessage",
        "CustomRdpProperty",
        "Tags",
        "HostPoolId",
    ],
    "AppGroups_Access": [
        "HostPoolName",
        "HostPoolType",
        "HostPoolResourceGroup",
        "ApplicationGroupName",
        "ApplicationGroupType",
        "ApplicationGroupFriendlyName",
        "Workspaces",
        "ApplicationName",
        "ApplicationFriendlyName",
        "ApplicationFilePath",
        "PrincipalType",
        "AssignedTo",
        "UserPrincipalName",
        "Mail",
        "Role",
        "AssignmentScope",
        "AssignmentLevel",
    ],
    "HostPools": [
        "HostPoolName",
        "HostPoolType",
        "HostPoolFriendlyName",
        "HostPoolDescription",
        "HostPoolResourceGroup",
        "Location",
        "LoadBalancerType",
        "MaxSessionLimit",
        "PreferredAppGroupType",
        "PersonalDesktopAssignmentType",
        "StartVMOnConnect",
        "ValidationEnvironment",
        "SessionHostCount",
        "ApplicationGroupCount",
        "CustomRdpProperty",
        "PublicNetworkAccess",
        "AgentUpdateType",
        "Tags",
        "HostPoolId",
    ],
    "SessionHosts": [
        "HostPoolName",
        "HostPoolType",
        "SessionHostName",
        "OSName",
        "OSVersion",
        "OSType",
        "SessionHostStatus",
        "SessionHostAllowNewSession",
        "Sessions",
        "AssignedUser",
        "LastHeartBeat",
        "AgentVersion",
        "SxSStackVersion",
        "UpdateState",
        "UpdateErrorMessage",
        "SessionHostResourceId",
    ],
    "ApplicationGroups": [
        "HostPoolName",
        "HostPoolType",
        "HostPoolResourceGroup",
        "ApplicationGroupName",
        "ApplicationGroupType",
        "ApplicationGroupFriendlyName",
        "ApplicationGroupDescription",
        "ApplicationGroupResourceGroup",
        "Location",
        "Workspaces",
        "ApplicationCount",
        "Applications",
        "AssignmentCount",
        "AssignedTo",
        "ApplicationGroupId",
    ],
    "Assignments": [
        "HostPoolName",
        "HostPoolType",
        "ApplicationGroupName",
        "ApplicationGroupType",
        "ApplicationGroupResourceGroup",
        "PrincipalType",
        "AssignedTo",
        "UserPrincipalName",
        "Mail",
        "PrincipalId",
        "Role",
        "AssignmentLevel",
        "AssignmentScope",
        "RoleAssignmentId",
    ],
    "Applications": [
        "HostPoolName",
        "HostPoolType",
        "ApplicationGroupName",
        "ApplicationGroupType",
        "ApplicationName",
        "ApplicationFriendlyName",
        "ApplicationType",
        "FilePath",
        "CommandLineSetting",
        "CommandLineArguments",
        "ShowInPortal",
        "IconPath",
        "AssignedTo",
    ],
    "Workspaces": [
        "WorkspaceName",
        "FriendlyName",
        "Description",
        "ResourceGroup",
        "Location",
        "ApplicationGroupName",
        "ApplicationGroupId",
        "WorkspaceId",
    ],
    "Summary": ["Metric", "Value"],
}


def write_workbook(path: str, data: dict[str, list[dict[str, Any]]]) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for sheet_name, columns in SHEET_COLUMNS.items():
        write_sheet(workbook, sheet_name, data.get(sheet_name, []), columns)
    workbook.save(path)
    LOGGER.info("Wrote %s", path)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export Azure Virtual Desktop host pools, session hosts, application groups, "
        "assignments, and applications to a multi-sheet Excel workbook."
    )
    parser.add_argument(
        "--subscription-id",
        required=True,
        help="Azure subscription ID to scan.",
    )
    parser.add_argument(
        "--tenant-id",
        default=None,
        help="Optional tenant ID (useful with device/interactive auth).",
    )
    parser.add_argument(
        "--auth",
        choices=("default", "cli", "device", "interactive"),
        default="default",
        help="Authentication method. default uses DefaultAzureCredential (az login works).",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="Output .xlsx path. Defaults to avd-inventory-<subscription>-<timestamp>.xlsx",
    )
    parser.add_argument(
        "--include-inherited-assignments",
        action="store_true",
        help="Include role assignments inherited from resource group or subscription.",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        help="Enable debug logging.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
    )
    if not args.verbose:
        logging.getLogger("azure").setLevel(logging.WARNING)
        logging.getLogger("azure.identity").setLevel(logging.WARNING)

    LOGGER.info(
        "AVD inventory exporter %s (%s)",
        globals().get("SCRIPT_VERSION", "unknown"),
        __file__,
    )

    credential = get_credential(args.auth, args.tenant_id)
    collector = AvdInventoryCollector(
        credential,
        args.subscription_id,
        include_inherited_assignments=args.include_inherited_assignments,
    )
    data = collector.collect()

    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    output = args.output or f"avd-inventory-{args.subscription_id}-{stamp}.xlsx"
    write_workbook(output, data)
    print(f"Inventory written to {output}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        LOGGER.error("Cancelled")
        raise SystemExit(130)
    except HttpResponseError as exc:
        LOGGER.error("Azure API error: %s", exc)
        raise SystemExit(1)
